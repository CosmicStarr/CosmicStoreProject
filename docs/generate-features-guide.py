"""Generate CosmicStore Features & User Guide Word document."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(r"C:\Users\Norma\CosmicStoreProject\docs\CosmicStore-Features-and-User-Guide.docx")
FALLBACK_OUT = Path(r"C:\Users\Norma\CosmicStoreProject\docs\_features-guide-regen.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()

# Narrow default font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def h(level, text):
    doc.add_heading(text, level=level)


def p(text, *, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    return para


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            t.rows[r].cells[c].text = value
    doc.add_paragraph()


# ---------- Title ----------
title = doc.add_heading("CosmicStore Features & User Guide", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run(
    f"Customer storefront, admin operations, CJ dropshipping, and Azure readiness\n"
    f"Generated {date.today().isoformat()}"
).italic = True

p(
    "CosmicStore is an Angular storefront plus ASP.NET Core API that sells products fulfilled through "
    "CJ Dropshipping. Customers pay with Stripe. CJ is paid from your wallet only after an order is paid. "
    "This guide covers every major feature, how to use it, and what was verified in testing."
)

# ---------- TOC-like overview ----------
h(1, "1. Application overview")
p("Stack and roles:")
bullets(
    [
        "Frontend: Angular app (CosmicStock) — shopper UI and Admin UI.",
        "Backend: CosmicStoreAPI (.NET) — Identity/JWT, catalog, cart, orders, Stripe, CJ.",
        "Data: SQL Server (store + CJ staging), Redis (cart/cache/sync timestamps).",
        "Roles: Guest shopper, registered customer, Admin.",
    ]
)

h(2, "1.1 Main areas")
table(
    ["Area", "URL examples", "Who uses it"],
    [
        ["Home & shop", "/ , /store , /store/{id}", "Everyone"],
        ["Cart & checkout", "/cart , /checkout", "Everyone (guest or signed in)"],
        ["Account", "/login , /register , /account/profile", "Customers"],
        ["Orders", "/orders , /order/manage", "Customers / guests"],
        ["Wishlist / registry", "/account/wishlist , /wishlist/{publicId}", "Customers / public buyers"],
        ["Legal", "/about , /shipping , /returns , /privacy , /terms", "Everyone"],
        ["Admin", "/admin/dashboard , /admin/orders , /admin/users , /admin/settings", "Admins only"],
    ],
)

# ---------- Storefront ----------
h(1, "2. Storefront (shopper features)")

h(2, "2.1 Home page")
p("Open the site root (/). You should see:")
bullets(
    [
        "Brand hero with Shop Now → opens /store.",
        "Category shortcuts that filter the catalog.",
        "Product tabs: New Arrivals, Top Selling, Featured.",
        "Links to shipping and returns policies.",
    ]
)
p("Expected: Clicking a product card opens product detail. New Arrival badges expire after 7 days (worker clears them).")

h(2, "2.2 Shop / catalog (/store)")
numbered(
    [
        "Open Shop from the navbar or Home → Shop Now.",
        "Use search (navbar or store filters), category chips, price filter, sort, and page size.",
        "Signed-in shoppers may see a “Bought before” section from recent purchases.",
        "Click a product to open /store/{productId}.",
    ]
)

h(2, "2.3 Product detail (/store/{id})")
numbered(
    [
        "Review gallery, description, and price.",
        "Choose a type/variant and quantity.",
        "Click Add to cart (updates cart and may open the side cart panel).",
        "Signed-in users can use the wishlist heart.",
        "Use share actions (native share / social / copy link) if desired.",
        "Related products appear below the main product.",
    ]
)

h(2, "2.4 Cart")
p("Two entry points:")
bullets(
    [
        "Cart icon → side panel: remove lines, View Full Cart, Checkout.",
        "Full page /cart: remove lines, Continue Shopping, Proceed to Checkout.",
    ]
)
p("Cart data lives in Redis; guest carts are keyed in browser storage and can merge after login.")

h(2, "2.5 Checkout & Stripe payment (/checkout)")
numbered(
    [
        "Enter contact email (order confirmation destination).",
        "Enter shipping address — or pick a saved address if signed in. Gift-registry carts lock shipping to the registry destination; only card ZIP is needed for Stripe AVS.",
        "Wait for shipping quotes (CJ freight). Select a shipping method.",
        "Enter card details in Stripe Elements (card data never hits CosmicStore servers).",
        "Accept Terms, Privacy, Shipping, and Returns (required).",
        "Click Pay. On success you go to /checkout/confirmation.",
    ]
)
p(
    "Server flow: create/update PaymentIntent → Stripe confirms card → create order → clear cart. "
    "CJ fulfillment is submitted only when your CJ wallet can cover product cost + freight. "
    "Otherwise the order stays paid until you top up and sync.",
    bold=False,
)
p("Confirmation page: shows order summary; guests get a Manage this order link from email.")

h(2, "2.6 Legal pages")
table(
    ["Page", "Path"],
    [
        ["About", "/about"],
        ["Shipping & Delivery", "/shipping"],
        ["Returns & Refunds", "/returns"],
        ["Privacy Policy", "/privacy"],
        ["Terms & Conditions", "/terms"],
    ],
)
p("Register and checkout require accepting the current legal version (must match frontend and API).")

# ---------- Account ----------
h(1, "3. Accounts and authentication")

h(2, "3.1 Register (/register)")
numbered(
    [
        "Enter username, email, password (complexity rules apply), confirm password.",
        "Check agreement to Terms, Privacy, Shipping, Returns.",
        "Submit Register. You receive a confirmation email.",
        "Emails listed in Store:AdminEmails are granted the Admin role automatically.",
    ]
)

h(2, "3.2 Login (/login)")
numbered(
    [
        "Enter email and password.",
        "Admins land on /admin/dashboard; others go to /store (or returnUrl).",
        "Unknown email can redirect to register.",
        "Use Forgot password for reset email → /account/reset-password.",
    ]
)

h(2, "3.3 Profile (/account/profile)")
p("Registered (non-guest) users can:")
bullets(
    [
        "Update username.",
        "Change password.",
        "Change email (confirm current password → verify new email via link).",
        "Manage saved shipping addresses (add / edit / delete / default).",
    ]
)

h(2, "3.4 Guest shopping")
bullets(
    [
        "Guests can browse, cart, and checkout without creating an account.",
        "Guest orders are managed via emailed link → /order/manage (email + ZIP verify).",
        "Guests cannot access Profile; wishlist requires sign-in.",
    ]
)

h(2, "3.5 Lock account")
p(
    "Email-change alerts can include a lock link (/account/lock-account). "
    "A locked user cannot sign in until an Admin unlocks them on /admin/users."
)

# ---------- Orders ----------
h(1, "4. Orders")

h(2, "4.1 Signed-in orders (/orders)")
numbered(
    [
        "Open Orders from the navbar or dashboard menu.",
        "List shows recent orders (last ~3 months).",
        "Open an order for tracking, shipping, and actions.",
        "Before ship: Cancel order or cancel a line item (Stripe refund path).",
        "After ship: Request / update a refund with reason; return tracking when applicable.",
    ]
)

h(2, "4.2 Guest order manage (/order/manage)")
numbered(
    [
        "Open the private link from the confirmation email (includes order id + token).",
        "Enter checkout email; enter ZIP for non-registry orders.",
        "Verify, then cancel / cancel item / request refund similar to signed-in users.",
    ]
)

# ---------- Wishlist ----------
h(1, "5. Wishlist and gift registry")

h(2, "5.1 Private wishlist (/account/wishlist)")
numbered(
    [
        "Sign in and open Wishlist.",
        "Add products from product detail (heart).",
        "Optionally enable Share as gift registry, pick a delivery address, save settings.",
        "Copy the public share link.",
        "Add items to cart or remove them.",
    ]
)

h(2, "5.2 Public registry (/wishlist/{publicId})")
numbered(
    [
        "Buyer opens the shared URL (no full street shown — masked destination).",
        "Add registry items to cart.",
        "Checkout ships to the registry address; buyer only enters card ZIP for payment.",
    ]
)

# ---------- Admin ----------
h(1, "6. Admin features")
p("Requires Admin role. Open /admin (sidebar: Products, Orders, Users, Settings).")

h(2, "6.1 Products dashboard (/admin/dashboard)")
numbered(
    [
        "Browse CJ staging catalog (search, category, sort, pagination).",
        "Publish to Store for one product, or Publish All on Page.",
        "Edit Details → /admin/edit-Product/{id}.",
        "Add Product → /admin/add-product.",
    ]
)
p("Publish multiplies CJ cost by the default markup from Settings (live DB value).")

h(2, "6.2 Create / edit product")
numbered(
    [
        "Optional: enter CJ pid or SKU → Load from CJ (fills name, price, images, types).",
        "Edit name, price, category, descriptions, flags (New Arrival / Top Selling / Featured).",
        "Manage types and images. Confirm before delete.",
        "Save (staging/admin catalog) and/or Publish to Store.",
        "Note: Mapping CJ variants for fulfillment does NOT auto-publish new options to the storefront. Edit and save/publish to show them.",
    ]
)

h(2, "6.3 Admin orders (/admin/orders)")
numbered(
    [
        "Review paid/shipped/refund status and tracking.",
        "FTC shipping compliance: paid orders still unfulfilled by CJ for 20+ days are sorted to the top, highlighted, and show an “Nd unfulfilled” flag (banner when any exist). Act before day 30 — US Mail Order Rule requires notify + refund offer if you cannot ship within 30 days.",
        "Sync all from CJ or sync one order.",
        "Cancel unshipped CJ orders; refund when policy allows.",
        "After customer return: open CJ dispute with return tracking; refresh until warehouse receipt; then refund.",
    ]
)

h(2, "6.4 Users (/admin/users)")
table(
    ["Action", "Behavior"],
    [
        ["Grant / Revoke Admin", "Cannot remove the last Admin"],
        ["Lock / Unlock", "Lockout; cannot lock yourself"],
        ["Reset password", "Emails the standard reset link"],
        ["Delete", "Cannot delete yourself or the last Admin"],
    ],
)
p("Columns: Email, Username, Roles, Confirmed, Locked.")

h(2, "6.5 Settings (/admin/settings)")
numbered(
    [
        "Set Default markup (e.g. 2.0×) and Catalog sync interval (Every 6 hours / Every 12 hours).",
        "Click Save settings — interval applies to the catalog worker without restart.",
        "View CJ wallet balance; Refresh balance. Top up when low so paid orders can fulfill.",
        "Force sync: Map product variants, Refresh stock, Sync order statuses.",
        "Read last-run timestamps for Catalog, Variants, Stock, Orders (or Never).",
    ]
)

# ---------- Workers ----------
h(1, "7. Background workers")
table(
    ["Worker", "What it does", "Cadence"],
    [
        [
            "CJ catalog sync (delta)",
            "First run: full category bootstrap into FlatProducts / FlatCategories. Later runs: listV2 timeStart/timeEnd from Redis LastRunTimestamp (cj:catalog:last-sync) so only newly listed products are upserted. Published storefront pids in that delta get variant + stock refresh. Cursor TTL is 90 days so the next interval can read it.",
            "6 or 12 hours (Admin Settings); does not skip when the Redis key exists — the key is the delta cursor",
        ],
        [
            "Order status worker (Production)",
            "Polls CJ shipment status; periodic full stock sync; expires New Arrivals older than 7 days",
            "Default every 30 minutes; stock every N polls",
        ],
    ],
)
p(
    "CJ listV2 has no true modifiedSince filter; delta uses listing-time windows. "
    "Warehouse stock for all mapped variants continues on the order-status worker cadence."
)

# ---------- Testing ----------
h(1, "8. Test results")
p(
    f"Testing window: {date.today().isoformat()}. Frontend http://127.0.0.1:4200 and API https://localhost:5001 were used where available."
)

h(2, "8.1 Automated / API checks (executed)")
table(
    ["Check", "Result", "Notes"],
    [
        ["GET /api/Products/joined-products", "PASS", "HTTP 200, product JSON returned"],
        ["GET /api/Products/categories", "PASS", "HTTP 200"],
        ["GET /api/Home/highlighted/NewArrival", "PASS", "HTTP 200"],
        ["GET /api/Payment/config", "PASS", "Publishable key returned"],
        ["GET /api/Account (no token)", "PASS (expected)", "HTTP 401 Unauthorized"],
        ["Angular app /", "PASS", "Dev server HTTP 200"],
        ["EF migrations up to date", "PASS", "Includes StoreRuntimeSettings + NewArrivalMarkedAt"],
    ],
)

h(2, "8.2 Browser UI checks (executed against http://127.0.0.1:4200)")
table(
    ["Check", "Result", "Notes"],
    [
        ["Home /", "PASS", "Hero, categories, New Arrivals / Top Selling / Featured product grids loaded"],
        ["Catalog /store", "PASS", "Price filter, search, sort, page size controls present"],
        ["Product detail", "PASS", "Gallery, qty, Add to cart, wishlist, share, description for gazebo SKU CJYL2466147"],
        ["Add to cart", "PASS", "Button showed Adding… then cart listed 1 item"],
        ["Cart /cart", "PASS", "Line item, Remove, Continue Shopping, Proceed to Checkout"],
        ["Checkout /checkout", "PASS", "Contact + shipping fields; CJPacket Ordinary $4.99 quote; Pay $804.98 disabled until form complete"],
        ["Login /login", "PASS", "Email/password, Forgot password, Register, guest shopping link"],
        ["Register /register", "PASS", "Fields + required legal checkbox (Terms/Privacy/Shipping/Returns)"],
        ["Forgot password", "PASS", "Reset your password form with email + submit"],
        ["Terms /terms", "PASS", "Version 2026-09-21.3 content rendered"],
        ["Shipping /shipping", "PASS", "Shipping & Delivery heading"],
        ["Guest manage /order/manage", "PASS", "Manage a guest order page loads"],
        ["Admin /admin/dashboard", "PASS (guard)", "Redirected to /login?returnUrl=/admin/dashboard when signed out"],
        ["Paid Stripe checkout", "NOT RUN", "Would charge card; see expected behavior below"],
        ["Admin Users/Settings while signed in", "NOT RUN", "No admin session in this test pass"],
        ["Email send (confirm/reset)", "NOT RUN", "Requires Graph mail; flow documented below"],
    ],
)

h(2, "8.3 Expected behavior (not fully exercised end-to-end)")
table(
    ["Flow", "What should happen"],
    [
        ["Full Stripe checkout", "Card confirms → order created → confirmation page + email; CJ submit if wallet covers cost"],
        ["Guest manage link", "Email link opens /order/manage with token; verify email+ZIP; cancel/refund actions work"],
        ["Admin publish", "Staging product appears on storefront after Publish; markup from Settings applies"],
        ["Variant sync", "CJ vids stored for fulfillment; new options stay off storefront until edit/save/publish"],
        ["Password reset / confirm email", "Graph mail sends HTML templates; links hit ReturnPath URLs"],
        ["Registry checkout", "Shipping locked to registry address; buyer pays Stripe; ship-to is recipient"],
        ["Admin refund after return", "Dispute opened with return tracking → CJ receipt → Stripe refund"],
        ["Catalog interval 6↔12", "Save Settings; worker delay re-reads hours next loop without API restart"],
    ],
)

h(2, "8.4 Known local caveats")
bullets(
    [
        "API must be running with Development secrets (SQL, Redis, JWT, Stripe test keys, Graph mail).",
        "Email features need Microsoft Graph signed in (dev script) or client-credentials in Production.",
        "Do not point live Stripe webhooks at localhost without a tunnel; use Stripe test mode locally.",
        "Admin routes require an Admin JWT; bootstrap via Store:AdminEmails on register.",
    ]
)

# ---------- Azure ----------
h(1, "9. Azure deployment readiness (next step)")
p(
    "You asked to deploy to Azure afterwards. Use this checklist when you are ready. "
    "Do not commit live secrets to git; set App Settings / Key Vault."
)

h(2, "9.1 Required production settings")
table(
    ["Setting", "Purpose"],
    [
        ["Store:PublicOrigin", "Live HTTPS storefront origin (emails, JWT issuer, CORS)"],
        ["Store:CorsOrigins", "Allowed browser origins if API is on another host"],
        ["Store:AdminEmails", "Emails that become Admin on register"],
        ["ConnectionStrings:DefaultConnection", "Azure SQL"],
        ["ConnectionStrings:RedisConnection", "Azure Cache for Redis"],
        ["JWT:SecretKey", "≥ 64 character secret"],
        ["Stripe:SecretKey / PublishableKey / WebhookSecret", "Live keys + webhook to /api/Payment/webhook"],
        ["Graph:ClientId / ClientSecret / TenantId", "App-only mail (AuthMode=ClientCredentials)"],
        ["ReturnPath:SenderEmail", "From address"],
        ["CJDropshipping:ApiKey", "CJ API access"],
    ],
)

h(2, "9.2 Suggested Azure layout")
bullets(
    [
        "App Service (API) + Static Web App or App Service (Angular) — or single host with Angular under wwwroot / CDN.",
        "Azure SQL Database + Azure Cache for Redis.",
        "Stripe Dashboard webhook → https://{api-host}/api/Payment/webhook.",
        "Set ASPNETCORE_ENVIRONMENT=Production so ProductionGuard validates settings.",
        "Optional: Database:MigrateOnStartup=true for first deploy, then turn off.",
        "Ensure Angular environment.prod.ts baseUrl points at /api/ or your API host.",
    ]
)

h(2, "9.3 Post-deploy smoke test")
numbered(
    [
        "Home and /store load over HTTPS.",
        "Register + confirmation email arrives.",
        "Admin Settings shows CJ wallet.",
        "Publish one product; it appears on the storefront.",
        "Place a small Stripe test/live order; confirmation email + order in Admin.",
        "Guest manage link works for a guest checkout.",
        "Force sync buttons update last-run times.",
    ]
)

h(1, "10. Quick route map")
p(
    "/ · /store · /store/:id · /cart · /checkout · /checkout/confirmation · /orders · /orders/:id · "
    "/order/manage · /account/wishlist · /wishlist/:publicId · /account/profile · /login · /register · "
    "/about · /shipping · /returns · /privacy · /terms · /admin/dashboard · /admin/add-product · "
    "/admin/edit-Product/:id · /admin/orders · /admin/orders/:orderId · /admin/users · /admin/settings"
)

p("End of document.", bold=True)

try:
    doc.save(OUT)
    print(f"Wrote {OUT}")
except PermissionError:
    doc.save(FALLBACK_OUT)
    print(f"Primary docx locked; wrote {FALLBACK_OUT}")
    print(f"Close Word and re-run, or rename {FALLBACK_OUT.name} over {OUT.name}.")
