"""Generate CosmicStore project outline Word document."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

title = doc.add_heading("CosmicStore Project Outline", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph("CJ Dropshipping E-Commerce Store")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(14)
subtitle.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

meta = doc.add_paragraph(
    "Updated 11 September 2026\n"
    "Stack: Angular 22 · .NET 10 Web API · Redis · SQL Server · EF Core\n"
    "Dual DbContext: CJ staging (dbo) + storefront/Identity (store)"
)
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.runs[0].font.size = Pt(10)
meta.runs[0].font.italic = True

doc.add_paragraph()

# ---------------------------------------------------------------------------
doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "CosmicStore is now a shoppable local store: customers can browse, add SKUs to a cart, "
    "pay with Stripe test keys, and create orders. Admin can publish the CJ staging catalog "
    "or add a product by CJ pid (variants/vids are saved). Email uses Microsoft Graph. "
    "CJ catalog sync workers and live CJ fulfillment (create order / pay wallet) are "
    "intentionally disabled so test checkouts do not place real shipments."
)
doc.add_paragraph(
    "What is left is mostly go-live work: turn fulfillment back on, keep catalog/stock in "
    "sync automatically, fix a few pid/vid mapping gaps, then Docker/CI, production config, "
    "and tests."
)

# ---------------------------------------------------------------------------
doc.add_heading("What's Left To Do", level=1)
p = doc.add_paragraph()
run = p.add_run("Do these before taking real customer orders.")
run.bold = True

doc.add_heading("Next (fulfillment & catalog)", level=2)
for item in [
    "Re-enable CJ createOrderV3 + payBalanceV2 in OrderService (currently commented out). Paid orders stay at PaymentRecevied and never submit to CJ.",
    "Re-enable CjOrderStatusWorker so shipment status/tracking updates after CJ accepts an order.",
    "Re-enable CJProductSyncWorker (or run it on demand) to refresh dbo.FlatProducts. Hosted service is commented out in Program.cs.",
    "Store the CJ pid on store.Products for Add Product (Guid ids). Variant re-sync today calls GetProductVariantsAsync(product.Id), which only works when Id is the pid.",
    "Stop writing Products.CjVariantId = product.Id on catalog publish/edit. That field should be a vid; ProductVariant.CjVariantId is the real mapping.",
    "Make admin Refund call Stripe (and CJ if needed). RefundOrderAsync only sets Status = Refunded locally.",
]:
    doc.add_paragraph(item, style="List Number")

doc.add_heading("Then (operations)", level=2)
for item in [
    "Admin role assignment: Register currently grants Admin only to a hardcoded email. Replace with a real invite/role UI.",
    "Optional: host product images instead of hotlinking CJ URLs.",
    "Optional: live warehouse stock on the storefront (stock sync API exists; it is not on a timer).",
    "Optional: storefront variant picker. Today the selected gallery photo's skuPhoto is the SKU that maps to a vid at checkout.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Last (production)", level=2)
for item in [
    "Docker Compose for API + Angular + Redis + SQL.",
    "CI/CD pipeline.",
    "Production CORS, HTTPS, health checks.",
    "Structured logging (Serilog) and error tracking (App Insights / Sentry).",
    "Automated tests — there are currently no .cs or .spec.ts test files.",
]:
    doc.add_paragraph(item, style="List Bullet")

# ---------------------------------------------------------------------------
doc.add_heading("What You Already Have", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Layer"
hdr[1].text = "Completed Features"
for cell in hdr:
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

rows = [
    (
        "CJ Integration",
        "Auth manager, listV2 sync worker (code present, hosted service off), freight quotes, "
        "wallet balance, variant/stock APIs, createOrder/payBalance implemented but not called",
    ),
    (
        "Catalog Pipeline",
        "FlatProducts staging → publish/bulk publish with markup; Add Product by pid "
        "(Guid store id + ProductVariant vids); storefront stored procs + Redis cache",
    ),
    (
        "Shopping",
        "Redis cart merged by SKU, guest/user merge, shipping quotes, Stripe PaymentIntent + webhook, "
        "orders with OrderItem.CjVariantId resolved from ProductVariant",
    ),
    (
        "Authentication",
        "Register/login JWT, Graph email confirm + forgot/reset password, profile, "
        "auth/admin/guest guards",
    ),
    (
        "Frontend",
        "Home, store (search/filter/sort/page), product detail + SKU from gallery, cart, checkout, "
        "orders, wishlist, addresses, admin dashboard/add/edit/orders/users/settings",
    ),
    (
        "Infrastructure",
        "Exception middleware, pagination headers, CORS, Redis, user secrets, Graph mail auth cache",
    ),
]
for layer, features in rows:
    row = table.add_row().cells
    row[0].text = layer
    row[1].text = features

doc.add_paragraph()

# ---------------------------------------------------------------------------
doc.add_heading("Architecture Overview", level=1)
doc.add_paragraph("Two-tier catalog plus a live storefront:")
for text in [
    "CJ API → CJProductSyncWorker → dbo.FlatProducts / FlatCategories (ApplicationDbContext). Worker currently off.",
    "Admin publish → store.Products (id = CJ pid) + ProductImages.",
    "Add Product → store.Products (new Guid) + ProductVariant rows (vid + SKU) from the pid you enter.",
    "Storefront API → stored procedures → ProductWithPictureDto grouped into ProductResponseDto → Angular (Redis cached).",
    "Cart is Redis JSON (ShoppingCart / CartItems). SQL only stores ShoppingCartSessionId.",
    "Checkout → Stripe → Order / OrderItem. Line SKU looks up ProductVariant.vid. CJ submit is skipped.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("DbContext Split", level=2)
ctx_table = doc.add_table(rows=1, cols=3)
ctx_table.style = "Table Grid"
h = ctx_table.rows[0].cells
h[0].text = "DbContext"
h[1].text = "Purpose"
h[2].text = "Key Entities"
for c in h:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
ctx_rows = [
    ("ApplicationDbContext", "CJ staging / admin catalog", "FlatProduct, FlatCategory"),
    (
        "ApplicationDbStoreContext",
        "Storefront + Identity",
        "Products, ProductImage, ProductVariant, AppUser, Order, OrderItem, WishlistItem, UserAddress, ShoppingCartSessionId",
    ),
]
for a, b, c in ctx_rows:
    r = ctx_table.add_row().cells
    r[0].text = a
    r[1].text = b
    r[2].text = c

doc.add_paragraph()
id_note = doc.add_paragraph()
id_note.add_run("IDs: ").bold = True
id_note.add_run(
    "FlatProducts.Id is a CJ pid. Catalog publishes keep that pid as store.Products.Id. "
    "Add Product uses a new Guid and stores vids on ProductVariant. Order lines store SKU + CjVariantId (vid)."
)

# ---------------------------------------------------------------------------
GREEN = RGBColor(0x16, 0xA3, 0x4A)
ORANGE = RGBColor(0xFF, 0x91, 0x00)
GRAY = RGBColor(0x64, 0x74, 0x8B)

phases = [
    (
        "Phase 0 — Foundation Fixes",
        "COMPLETE",
        [
            "DI: CJ, cart, Redis, JWT options",
            "JWT issuer/audience and Redis connection string",
            "Login UI, routes, product links, edit form population",
            "Secrets moved to user secrets (JWT, CJ, Graph, Stripe)",
        ],
    ),
    (
        "Phase 1 — Core Shopping Flow",
        "COMPLETE (fulfillment off)",
        [
            "CartController, guest/user merge, Angular cart panel + page",
            "Checkout, Order/OrderItem persistence, Stripe test payments + webhook",
            "User order history and order detail",
            "SKU → ProductVariant vid on order create (CJ submit commented out)",
        ],
    ),
    (
        "Phase 2 — Catalog & Admin Pipeline",
        "COMPLETE",
        [
            "Publish / bulk publish FlatProduct → store.Products with markup",
            "Edit updates existing store products; Add Product creates a Guid product",
            "Add Product: enter CJ pid → Load from CJ → save variants (vid + SKU + images)",
            "Admin layout, auth+admin guards, dashboard, orders, users, settings",
            "Stored procedures still used for storefront lists",
        ],
    ),
    (
        "Phase 3 — Auth & User Features",
        "COMPLETE",
        [
            "Microsoft Graph email (SMTP was blocked); confirm + forgot/reset password pages",
            "authGuard, adminGuard, guestGuard",
            "Profile, wishlist, saved addresses",
        ],
    ),
    (
        "Phase 4 — Storefront Polish",
        "COMPLETE",
        [
            "Search, category filter, price range, sort, pagination wired to API",
            "Related products; gallery skuPhoto shown and sent to cart",
            "Home/nav/footer use live catalog data; 404 and product-not-found pages",
            "Stock on Products is a store field (live CJ poll is Phase 5)",
        ],
    ),
    (
        "Phase 5 — CJ Dropshipping Operations",
        "PARTIAL — NEXT",
        [
            "DONE: ProductVariant table, freight quotes, wallet endpoint, pid import, SKU→vid at checkout",
            "OFF: createOrderV3 / payBalanceV2, order-status worker, catalog sync worker",
            "GAP: no CjProductId column on Products; catalog sync treats Products.Id as pid",
            "GAP: refund is local status only; returns flow not built",
        ],
    ),
    (
        "Phase 6 — Infrastructure & Production",
        "NOT STARTED",
        [
            "No Docker, CI/CD, health checks, or structured logging",
            "No unit, component, or E2E tests",
            "Production CORS / HTTPS still to configure",
        ],
    ),
]

for phase_name, status, items in phases:
    doc.add_heading(phase_name, level=1)
    p = doc.add_paragraph()
    run = p.add_run(f"Status: {status}")
    run.bold = True
    if status.startswith("COMPLETE"):
        run.font.color.rgb = GREEN
    elif "NEXT" in status or status.startswith("PARTIAL") or status.startswith("MOSTLY"):
        run.font.color.rgb = ORANGE
    else:
        run.font.color.rgb = GRAY
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph()

# ---------------------------------------------------------------------------
doc.add_heading("Suggested Priority Order", level=1)
pri = doc.add_table(rows=1, cols=3)
pri.style = "Table Grid"
ph = pri.rows[0].cells
ph[0].text = "Priority"
ph[1].text = "Work"
ph[2].text = "Why"
for c in ph:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
priorities = [
    ("P0", "Foundation (DI, JWT, Redis, login)", "DONE"),
    ("P1", "Cart + checkout + Stripe + orders", "DONE (local/test)"),
    ("P2", "Admin publish + Add Product by pid", "DONE"),
    ("P3", "Email + guards + profile/wishlist/addresses", "DONE"),
    ("P4", "Turn CJ fulfillment back on + status worker", "Needed for real shipments"),
    ("P5", "Persist pid on Guid products; fix CjVariantId misuse", "Reliable vid mapping"),
    ("P6", "Re-enable catalog/stock sync; real admin roles", "Keep catalog current"),
    ("P7", "Docker, CI, logging, tests, production CORS", "Go live"),
]
for a, b, c in priorities:
    r = pri.add_row().cells
    r[0].text = a
    r[1].text = b
    r[2].text = c

doc.add_paragraph()

# ---------------------------------------------------------------------------
doc.add_heading("Quick Reference: Done vs Left", level=1)
ref = doc.add_table(rows=1, cols=2)
ref.style = "Table Grid"
rh = ref.rows[0].cells
rh[0].text = "Done"
rh[1].text = "Still left"
for c in rh:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
refs = [
    ("Cart API + UI, SKU merge", "—"),
    ("Stripe test checkout + webhook", "Live Stripe keys / production webhook URL"),
    ("Orders with vid on each line", "Uncomment CJ create/pay; status worker"),
    ("Publish + Add Product pid import", "Store pid on Guid products for later re-sync"),
    ("ProductVariant SKU ↔ vid", "Do not copy pid into Products.CjVariantId"),
    ("Graph email confirm/reset", "Replace hardcoded Admin email"),
    ("Admin pages + guards", "Real Stripe refund; returns UI"),
    ("CJ freight quotes + wallet API", "Hosted catalog/stock workers"),
    ("Controller + repo XML comments", "Docker, CI, tests, logging, health checks"),
]
for a, b in refs:
    r = ref.add_row().cells
    r[0].text = a
    r[1].text = b

doc.add_paragraph()

# ---------------------------------------------------------------------------
doc.add_heading("Current API Endpoints", level=1)
api = doc.add_table(rows=1, cols=3)
api.style = "Table Grid"
ah = api.rows[0].cells
ah[0].text = "Controller"
ah[1].text = "Role"
ah[2].text = "Auth"
for c in ah:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
apis = [
    ("AccountController", "Register, login, current user, confirm email, forgot/reset password, profile, change password", "Mixed"),
    ("ProductsController", "Catalog list, categories, product detail, related", "None"),
    ("HomeController", "Highlighted strips, product by id", "None"),
    ("CartController", "Get/add/update/remove/merge/clear Redis cart", "Merge requires sign-in"),
    ("WishlistController", "CRUD wishlist", "Authorize"),
    ("UserAddressesController", "Saved checkout addresses", "Authorize"),
    ("ShippingController", "CJ freight quote from line SKUs", "None"),
    ("PaymentController", "PaymentIntent, Stripe webhook, publishable key", "Intent: Authorize"),
    ("OrdersController", "Checkout, my orders, order by id", "Authorize"),
    ("EditProductsController", "Staging grid, create, edit, publish, bulk, settings", "Admin"),
    ("AdminCjController", "Pid preview, wallet, variant/stock sync, cancel/refund", "Admin"),
    ("AdminOrdersController", "All orders", "Admin"),
    ("AdminUsersController", "User list", "Admin"),
    ("ErrorController", "JSON status-code body", "None"),
]
for a, b, c in apis:
    r = api.add_row().cells
    r[0].text = a
    r[1].text = b
    r[2].text = c

doc.add_paragraph()

# ---------------------------------------------------------------------------
doc.add_heading("Current Angular Routes", level=1)
routes = doc.add_table(rows=1, cols=3)
routes.style = "Table Grid"
rt = routes.rows[0].cells
rt[0].text = "Route"
rt[1].text = "Component"
rt[2].text = "Guard"
for c in rt:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
route_rows = [
    ("/", "HomeComponent", "None"),
    ("/store", "ProductsComponent", "None"),
    ("/store/:id", "ProductDetailComponent", "None"),
    ("/cart", "CartComponent", "None"),
    ("/checkout", "CheckoutComponent", "authGuard"),
    ("/orders, /orders/:id", "Orders / OrderDetail", "authGuard"),
    ("/login, /register", "Login / Register", "guestGuard"),
    ("/account/confirm-email", "ConfirmEmailComponent", "None"),
    ("/account/forgot-password, /reset-password", "Forgot / Reset", "guestGuard"),
    ("/account/profile, /wishlist", "Profile / Wishlist", "authGuard"),
    ("/admin/*", "Dashboard, add/edit product, orders, users, settings", "authGuard + adminGuard"),
]
for a, b, c in route_rows:
    r = routes.add_row().cells
    r[0].text = a
    r[1].text = b
    r[2].text = c

doc.add_paragraph()
footer = doc.add_paragraph(
    "CosmicStoreProject · outline regenerated 11 September 2026 · source: docs/generate-outline.py"
)
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.runs[0].font.size = Pt(9)
footer.runs[0].font.italic = True
footer.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

output_path = r"C:\Users\Norma\CosmicStoreProject\docs\CosmicStore-Project-Outline.docx"
try:
    doc.save(output_path)
except PermissionError:
    output_path = r"C:\Users\Norma\CosmicStoreProject\docs\CosmicStore-Project-Outline-updated.docx"
    doc.save(output_path)
    print("Original outline is open in another program; wrote a new copy instead.")
print(f"Created: {output_path}")
