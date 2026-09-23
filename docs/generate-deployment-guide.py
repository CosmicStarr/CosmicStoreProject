from docx import Document
from docx.shared import Pt
from datetime import date
from pathlib import Path

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def h1(text: str) -> None:
    doc.add_heading(text, level=1)


def h2(text: str) -> None:
    doc.add_heading(text, level=2)


def h3(text: str) -> None:
    doc.add_heading(text, level=3)


def para(text: str = "") -> None:
    doc.add_paragraph(text)


def bullet(text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def numbered(text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def code_block(text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def note(text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run("Note: ")
    run.bold = True
    p.add_run(text)


doc.add_heading("CosmicStore Azure Deployment Guide", 0)
sub = doc.add_paragraph()
r = sub.add_run(
    "Detailed record of steps taken to deploy the CosmicStore API and CosmicStock "
    "Angular storefront to Microsoft Azure"
)
r.italic = True

meta = doc.add_paragraph()
meta.add_run(f"Document date: {date.today().isoformat()}\n").bold = True
meta.add_run("Repository: ").bold = True
meta.add_run("CosmicStarr/CosmicStoreProject\n")
meta.add_run("Resource group: ").bold = True
meta.add_run("cosmicsolutions\n")

para(
    "This document describes the end-to-end Azure go-live for CosmicStore: App Service API, "
    "Azure SQL, Redis, GitHub Actions CI/CD, Static Web Apps for Angular, secrets handling, "
    "networking/firewalls, schema alignment, troubleshooting, and verification. "
    "Secrets are referenced by configuration name only—never paste live passwords or API keys into this file."
)

h1("1. Architecture overview")
para("Production uses two Azure front doors plus shared data services:")
bullet("Azure App Service (Linux) cosmicstoreapi — hosts CosmicStoreAPI.dll")
bullet("Azure Static Web Apps (Free) cosmicstock — hosts CosmicStock Angular SPA")
bullet("Azure SQL Database cosmicstoreap-database on cosmicstoreap-server")
bullet("Azure Cache for Redis cosmicstoreap-cache (SSL port 6380)")
bullet("GitHub Actions on branch main — builds and deploys API and frontend")

h2("1.1 Live URLs (as of go-live)")
bullet("API: https://cosmicstoreapi-hqgvbza3abfngkb7.westus3-01.azurewebsites.net")
bullet("Storefront: https://yellow-water-0c144b71e.6.azurestaticapps.net")
bullet(
    "Angular production baseUrl: "
    "https://cosmicstoreapi-hqgvbza3abfngkb7.westus3-01.azurewebsites.net/api/"
)
note(
    "Static Web Apps assigns a random default hostname (for example yellow-water-…). "
    "That hostname cannot be renamed to cosmicstore-…. Use a custom domain for branding."
)

h2("1.2 Solution components")
bullet("CosmicStoreAPI — ASP.NET Core Web API (.NET 10)")
bullet("Data / Models — EF Core, migrations, Redis helpers")
bullet("CosmicStock — Angular storefront (dist/CosmicStock/browser)")
bullet(
    "Scripts/store-procedures.sql — historical reference only; EF migrations are source of truth on Azure"
)

h1("2. Prerequisites and accounts")
numbered("Azure subscription with rights to App Service, Static Web Apps, SQL, and Redis.")
numbered("GitHub repo CosmicStarr/CosmicStoreProject with Actions enabled.")
numbered(
    "OIDC federated credentials for App Service deploy "
    "(AZUREAPPSERVICE_CLIENTID_*, TENANTID_*, SUBSCRIPTIONID_* secrets)."
)
numbered("Local secrets via dotnet user-secrets on CosmicStoreAPI (never committed).")
numbered("Node.js 22+ / npm for Angular; .NET 10 SDK for the API.")
numbered("Azure Portal access for firewalls, App Settings, and Static Web App management.")

h2("2.1 Secrets inventory (names only)")
para(
    "These keys must exist in Azure App Service Application settings (double-underscore form) "
    "and/or GitHub Actions secrets. Values come from a secure store—not git."
)

h3("App Service application settings (Production)")
bullet("ASPNETCORE_ENVIRONMENT = Production")
bullet("ConnectionStrings__DefaultConnection — Azure SQL")
bullet(
    "ConnectionStrings__RedisConnection — host:6380,password=…,ssl=True,abortConnect=False"
)
bullet("JWT__SecretKey, JWT__ValidIssuer, JWT__ValidAudience")
bullet("Stripe__PublishableKey, Stripe__SecretKey, Stripe__WebhookSecret, Stripe__AllowTestKeys")
bullet("Graph__AuthMode, Graph__ClientId, Graph__ClientSecret, Graph__TenantId")
bullet("CJDropshipping__ApiKey")
bullet("SMTP__Password (and related mail settings if used)")
bullet("Store__PublicOrigin — storefront HTTPS origin (Static Web App URL)")
bullet("Store__CorsOrigins__0 — storefront origin (use __1 for additional origins)")
bullet("Store__AdminEmails__0 — admin notification mailbox")
bullet("Store__AllowLocalInfrastructure = false")
bullet(
    "Database__MigrateOnStartup = true for first go-live; consider false after schema is stable"
)
bullet("ReturnPath__* — email deep-link paths")

h3("GitHub Actions secrets")
bullet("AZUREAPPSERVICE_CLIENTID_… / TENANTID_… / SUBSCRIPTIONID_… — OIDC for API deploy")
bullet(
    "AZURE_STATIC_WEB_APPS_API_TOKEN_COSMICSTOCK — Static Web Apps deployment token for Angular CI/CD"
)

h1("3. Backend API — Azure App Service")

h2("3.1 Identify the Web App")
para(
    "Production API: Linux App Service cosmicstoreapi in resource group cosmicsolutions, "
    "West US 3, hostname cosmicstoreapi-hqgvbza3abfngkb7.westus3-01.azurewebsites.net."
)

h2("3.2 Startup Command")
para(
    "Portal → cosmicstoreapi → Configuration → General settings. Set Startup Command to:"
)
code_block("dotnet CosmicStoreAPI.dll")
para(
    "Without this, Linux App Service may show the Azure welcome page after a successful deploy."
)

h2("3.3 GitHub Actions workflow (API)")
para("Workflow file: .github/workflows/main_cosmicstoreapi.yml")
para("Triggers: push to main, or workflow_dispatch.")

h3("Build job")
numbered("actions/checkout@v4")
numbered("actions/setup-dotnet@v4 with dotnet-version 10.x")
numbered("dotnet build CosmicStoreAPI/CosmicStoreAPI.csproj -c Release")
numbered(
    "dotnet publish CosmicStoreAPI/CosmicStoreAPI.csproj -c Release "
    "-o ${{ github.workspace }}/artifacts/api"
)
numbered("actions/upload-artifact@v4 name .net-app from the publish folder")

h3("Deploy job")
numbered(
    "actions/download-artifact@v4 name .net-app with path: .  (critical — see section 3.4)"
)
numbered("azure/login@v2 using OIDC client/tenant/subscription secrets")
numbered(
    "azure/webapps-deploy@v3 app-name cosmicstoreapi, slot Production, package: ."
)

h2("3.4 Critical fix: download-artifact nesting")
para(
    "download-artifact@v4 nests files under the artifact name when path is wrong. "
    "If package: . only contains a nested .net-app/ folder, CosmicStoreAPI.dll is not at wwwroot. "
    "Symptom: green Actions run + Azure welcome page."
)
para("Fix applied:")
code_block(
    "- name: Download artifact from build job\n"
    "  uses: actions/download-artifact@v4\n"
    "  with:\n"
    "    name: .net-app\n"
    "    path: .   # flatten so CosmicStoreAPI.dll is at package root"
)
para(
    "Publish into a workspace folder such as artifacts/api — never into ${{ env.DOTNET_ROOT }}/myapp."
)

h2("3.5 Ensure main contains the full application")
para(
    "Go-live failure mode: origin/main only had early scaffold + deploy workflow commits, "
    "while feature branches held the full API. Azure SQL had ProductTypes / TypeName schema, "
    "but the deployed binary still expected older FromSql shapes (Type column). "
    "Result: HTTP 500 on catalog endpoints — SqlException Invalid column name 'Type'."
)
para("Resolution:")
numbered("Merge the full application onto main (API, Data, Models, CosmicStock as needed).")
numbered("Push to main to rebuild and redeploy the current codebase.")
numbered("Confirm Actions build + deploy jobs succeed.")

h2("3.6 Application settings restore")
para(
    "Settings were mapped from local dotnet user-secrets to Azure __ keys. "
    "Caution: ARM GET of appsettings without /list can return empty properties; "
    "a PUT would wipe secrets. Always use:"
)
code_block(
    "POST .../providers/Microsoft.Web/sites/{site}/config/appsettings/list?api-version=2023-12-01"
)
para(
    "After restore, verify ~29 production keys (ConnectionStrings, JWT, Stripe, Graph, Store, "
    "Database__MigrateOnStartup, etc.), then restart the App Service."
)

h1("4. Data stores — Azure SQL and Redis")

h2("4.1 Azure SQL Database")
bullet("Server: cosmicstoreap-server.database.windows.net")
bullet("Database: cosmicstoreap-database")
bullet(
    "Store objects: store.GetProducts, store.Pictures, store.ProductTypes, "
    "SPs GetAllProductsWithPictures, GetSingleProductWithPictures, GetHighlightedProducts"
)

h3("Firewall")
numbered("Allow Azure services so App Service can reach SQL.")
numbered("Add developer client IPs for local tools against Azure SQL.")
numbered(
    "Optionally add App Service outbound IPs from App Service → Properties → Outbound IP addresses."
)

h3("Migrations")
para(
    "With Database__MigrateOnStartup=true, the API runs EF Core MigrateAsync at startup. "
    "History table: dbo.__EFMigrationsHistory."
)
para("Schema evolution relevant to go-live:")
bullet("PictureType originally added Pictures.Type and SP columns selecting pi.Type.")
bullet(
    "ProductTypeEntity / ProductTypePrice introduced ProductTypes, Pictures.ProductTypeId, "
    "dropped Pictures.Type, updated SPs to TypeName / TypeSku / TypePrice."
)
bullet(
    "Do not re-run older Scripts/store-procedures.sql on Azure after ProductTypes shipped — "
    "it re-adds Type and regresses the schema."
)

h2("4.2 Azure Cache for Redis")
bullet("Host: cosmicstoreap-cache.redis.cache.windows.net:6380")
bullet("Requires ssl=True")
bullet(
    "abortConnect=False lets the API start if Redis is briefly down; treat cache as best-effort"
)

h3("Firewall")
para(
    "Allow developer IPs and every App Service outbound / possibleOutbound IP. "
    "Missing outbound IPs caused catalog 500s (Redis timeout) while /api/Payment/config "
    "still returned 200 (no Redis)."
)
para("During go-live, missing outbound IPs were added after comparing:")
code_block(
    "site.properties.outboundIpAddresses\n"
    "site.properties.possibleOutboundIpAddresses"
)

h3("Cache resilience")
para(
    "CacheService catches RedisException on get/set/delete/increment and falls back to SQL / no-op "
    "so a Redis firewall miss does not take down the catalog."
)

h1("5. Frontend — Azure Static Web Apps")

h2("5.1 Create the Static Web App")
numbered("Resource name: cosmicstock")
numbered("Resource group: cosmicsolutions")
numbered("Region: West US 2")
numbered("SKU: Free")
numbered("Provider: None (token-based deploy)")
numbered(
    "Get deployment token: Portal → Manage deployment token, or ARM listSecrets"
)

h2("5.2 SPA configuration")
para(
    "File: CosmicStock/public/staticwebapp.config.json "
    "(copied into the Angular browser build). "
    "navigationFallback rewrites to /index.html so routes like /store do not 404 on refresh."
)

h2("5.3 Production environment")
para("CosmicStock/src/env/environment.prod.ts:")
code_block(
    "baseUrl: 'https://cosmicstoreapi-hqgvbza3abfngkb7.westus3-01.azurewebsites.net/api/'"
)

h2("5.4 First content deploy (CLI)")
numbered("From CosmicStock: npm run build:prod")
numbered(
    "Confirm output at CosmicStock/dist/CosmicStock/browser "
    "(index.html + staticwebapp.config.json)"
)
numbered("Deploy:")
code_block(
    "npx @azure/static-web-apps-cli deploy ./dist/CosmicStock/browser "
    "--deployment-token <TOKEN> --env production"
)
para(
    "Successful deploy serves https://yellow-water-0c144b71e.6.azurestaticapps.net"
)

h2("5.5 GitHub Actions workflow (frontend)")
para("Workflow file: .github/workflows/azure-static-web-apps-cosmicstock.yml")
bullet("Triggers on CosmicStock/** or workflow file changes to main (also workflow_dispatch)")
bullet("Node 22, npm ci, npm run build:prod")
bullet(
    "Azure/static-web-apps-deploy@v1 with skip_app_build: true, "
    "output_location: dist/CosmicStock/browser"
)
bullet("Requires secret AZURE_STATIC_WEB_APPS_API_TOKEN_COSMICSTOCK")
note(
    "Until that GitHub secret is set, the SWA workflow fails even if the site is already live "
    "from the CLI deploy."
)

h2("5.6 CORS and PublicOrigin")
para("After the SWA hostname is known, set App Service:")
bullet("Store__PublicOrigin = https://yellow-water-0c144b71e.6.azurestaticapps.net")
bullet("Store__CorsOrigins__0 = https://yellow-water-0c144b71e.6.azurestaticapps.net")
bullet("Optional Store__CorsOrigins__1 for additional origins")
para(
    "Mirror defaults in CosmicStoreAPI/appsettings.Production.json "
    "(App Settings override at runtime). Restart the API after CORS changes."
)

h1("6. Verification checklist")

h2("6.1 API")
bullet("GET /api/Payment/config → 200 with publishableKey")
bullet("GET /api/Products/categories → 200 (may be [] until products are published)")
bullet("GET /api/Products/joined-products → 200")
bullet("GET /api/home/highlighted/featured → 200")
bullet("Root / → API JSON 404 is OK (not Azure welcome HTML)")

h2("6.2 Storefront")
bullet("GET storefront / → 200 Angular shell")
bullet("GET /store → 200 (SPA fallback)")
bullet("Browser Network: XHR to …/api/… succeeds (CORS OK)")

h2("6.3 Database")
bullet("SELECT COUNT(*) FROM store.GetProducts — publish via admin when ready")
bullet("SPs return TypeName/TypeSku/TypePrice (not pi.Type) after ProductType migrations")

h2("6.4 CI/CD")
bullet("Push to main triggers API workflow; CosmicStock path changes trigger SWA workflow")
bullet("Both jobs green in GitHub Actions")

h1("7. Troubleshooting log (issues and fixes)")

h2("7.1 Welcome page after green deploy")
bullet("Cause: nested artifact path and/or empty Startup Command")
bullet("Fix: download-artifact path: . ; Startup Command = dotnet CosmicStoreAPI.dll")

h2("7.2 App settings wiped")
bullet("Cause: ARM GET appsettings (without /list) then PUT empty properties")
bullet("Fix: restore via Advanced edit or POST …/appsettings/list then PUT full map")

h2("7.3 Catalog 500 — Invalid column name Type")
bullet("Cause: old API binary vs new DB schema, or obsolete SP script")
bullet("Fix: deploy current main; do not apply obsolete store-procedures.sql")

h2("7.4 Catalog 500 — Redis timeout; Payment still 200")
bullet("Cause: Redis firewall missing App Service outbound IPs")
bullet("Fix: add outbound + possibleOutbound IPs; harden CacheService")

h2("7.5 Empty categories / products arrays")
bullet("Cause: store.GetProducts has zero rows")
bullet("Fix: admin publish / CJ sync — not an infrastructure failure")

h2("7.6 Cannot rename yellow-water-… hostname")
bullet("Cause: Static Web Apps random default hostname is immutable")
bullet("Fix: keep random URL or attach a custom domain")

h1("8. Ongoing operations")

h2("8.1 Redeploy API")
numbered("Merge to main (or workflow_dispatch on main_cosmicstoreapi.yml)")
numbered("Confirm deploy success")
numbered("Smoke-test Payment/config and Products/categories")

h2("8.2 Redeploy storefront")
numbered("Ensure AZURE_STATIC_WEB_APPS_API_TOKEN_COSMICSTOCK is set")
numbered("Push CosmicStock changes to main, or workflow_dispatch")
numbered("Or local: npm run build:prod && swa deploy …")

h2("8.3 Rotate secrets")
para(
    "Rotate any connection strings, API keys, or deployment tokens that appeared in chat logs, "
    "temp files, or clipboards during go-live; update Azure and GitHub; restart App Service."
)

h2("8.4 Optional: turn off MigrateOnStartup")
para(
    "After production schema is stable, set Database__MigrateOnStartup=false so routine restarts "
    "do not run migrations. Apply future schema changes deliberately."
)

h1("9. Appendix — key file map")
bullet(".github/workflows/main_cosmicstoreapi.yml — API CI/CD")
bullet(".github/workflows/azure-static-web-apps-cosmicstock.yml — Angular CI/CD")
bullet("CosmicStock/public/staticwebapp.config.json — SPA fallback")
bullet("CosmicStock/src/env/environment.prod.ts — production API baseUrl")
bullet("CosmicStoreAPI/appsettings.Production.json — non-secret production defaults")
bullet("CosmicStoreAPI/Program.cs — Redis, MigrateOnStartup, CORS")
bullet("Data/Util/StoreUrls.cs — PublicOrigin / CorsOrigins")
bullet("Data/Classes/CacheService.cs — Redis cache with failure tolerance")
bullet("Data/Migrations/* — EF schema including ProductTypes / SP updates")

h1("10. Appendix — App Service setting naming")
para(
    "ASP.NET Core on Azure maps hierarchy with double underscores. "
    "Example: ConnectionStrings:RedisConnection ↔ ConnectionStrings__RedisConnection."
)

end = doc.add_paragraph()
end.add_run("End of document.").italic = True

out = Path(r"C:\Users\Norma\CosmicStoreProject\docs\CosmicStore-Azure-Deployment-Guide.docx")
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
print("bytes", out.stat().st_size)
